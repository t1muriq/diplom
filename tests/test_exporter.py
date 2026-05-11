"""
Тесты модуля выгрузок XLSX.

Проверяют, что генерируемые файлы открываются как корректные Excel,
содержат правильные заголовки и данные, имеют форматирование
(заливку для оценок, жирные шапки).

Запуск: pytest tests/test_exporter.py -v
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "web_app"))

import pytest
from openpyxl import load_workbook

import exporter


def _load(buf):
    """Загружает workbook из BytesIO."""
    buf.seek(0)
    return load_workbook(buf)


def test_export_teacher_schedule_basic():
    rows = [
        {"lesson_date": date(2026, 4, 6), "lesson_number": 1,
         "class_name": "9А", "subject_name": "Математика",
         "room_number": "101", "building": "Основной"},
        {"lesson_date": date(2026, 4, 6), "lesson_number": 3,
         "class_name": "10А", "subject_name": "Алгебра",
         "room_number": "102", "building": "Основной"},
    ]
    buf = exporter.export_teacher_schedule(
        teacher_full_name="Сидоров А. В.",
        rows=rows,
        date_from=date(2026, 4, 6), date_to=date(2026, 4, 12),
    )
    wb = _load(buf)
    ws = wb.active
    assert "Сидоров" in ws["A1"].value
    assert ws.cell(row=4, column=1).value == "Дата"
    assert ws.cell(row=5, column=4).value == "9А"  # первый класс
    assert ws.cell(row=5, column=1).font.size  # шрифт задан


def test_export_class_grades_with_colors():
    students = [
        {"student_id": 1, "full_name": "Иванов Иван"},
        {"student_id": 2, "full_name": "Петрова Анна"},
    ]
    grade_dates = [date(2026, 4, 1), date(2026, 4, 8)]
    grades_map = {
        (1, date(2026, 4, 1)): {"value": 5, "comment": "тест"},
        (1, date(2026, 4, 8)): {"value": 4, "comment": ""},
        (2, date(2026, 4, 1)): {"value": 3, "comment": ""},
    }
    buf = exporter.export_class_grades(
        class_name="9А", subject_name="Математика",
        teacher_full_name="Сидоров А. В.",
        students=students, grade_dates=grade_dates, grades_map=grades_map,
    )
    wb = _load(buf)
    ws = wb.active

    assert "9А" in ws["A1"].value
    assert "Математика" in ws["A1"].value
    # Шапка содержит даты
    assert ws.cell(row=4, column=3).value == "01.04"
    # Первый ученик, оценка 5 — зелёная заливка
    cell_5 = ws.cell(row=5, column=3)
    assert cell_5.value == 5
    assert "C6EFCE" in (cell_5.fill.start_color.rgb or "")
    # Тройка — жёлтая
    cell_3 = ws.cell(row=6, column=3)
    assert cell_3.value == 3
    assert "FFF2CC" in (cell_3.fill.start_color.rgb or "")


def test_export_students_list_basic():
    rows = [
        {"last_name": "Иванов", "first_name": "Иван",
         "middle_name": "Иванович",
         "birth_date": date(2010, 5, 17),
         "admission_year": 2022, "student_status": "active",
         "class_name": "9А"},
    ]
    buf = exporter.export_students_list(rows)
    wb = _load(buf)
    ws = wb.active
    assert "Школа №1416" in ws["A1"].value
    assert ws.cell(row=4, column=2).value == "Иванов"


def test_export_log_basic():
    rows = [
        (1, "students.xlsx", "2026-04-15 10:00:00", 12, "completed"),
        (2, "bad.csv", "2026-04-15 11:00:00", 0, "failed"),
    ]
    buf = exporter.export_log(
        title="Журнал импорта",
        headers=["ID", "Источник", "Время", "Загружено", "Статус"],
        rows=rows,
    )
    wb = _load(buf)
    ws = wb.active
    assert ws["A1"].value == "Журнал импорта"
    assert ws.cell(row=3, column=1).value == "ID"
    assert ws.cell(row=4, column=2).value == "students.xlsx"
    assert ws.cell(row=5, column=5).value == "failed"


def test_empty_data_does_not_crash():
    """Пустые данные не должны ломать генерацию."""
    buf = exporter.export_teacher_schedule(
        teacher_full_name="Иван Иванов",
        rows=[],
        date_from=date(2026, 4, 1), date_to=date(2026, 4, 7),
    )
    wb = _load(buf)
    ws = wb.active
    # Файл всё равно создан, просто содержит сообщение
    assert "Нет занятий" in ws.cell(row=5, column=1).value
